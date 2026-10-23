#!/usr/bin/env python3
"""List the functional annotations shared by many genes, flagging the transposon-like ones, for manual screening.

NOT part of the pipeline: an analysis tool (e.g. choosing which TE-like functions to drop from the final
annotation). Each term is counted once per gene, on the main protein. Sources:

  InterPro     InterProScan entry (IPR accession and name)
  Pfam         InterProScan Pfam signature
  PANTHER      InterProScan PANTHER family / subfamily
  eggNOG       eggNOG-mapper Description
  Diamond2GO   description of the Diamond2GO hit (accession removed)

Terms with more than --min-genes genes are kept. te_keyword gives the first transposon keyword found in the
term name (classify_gene_functions.py list, extended); it is a hint for
screening, not a decision. With --new-genes, the share of genes absent from the previous annotation is added.

Outputs: <out>.tsv and <out>.xlsx (one sheet per source plus an "all" sheet and a README).
"""
import argparse
import collections
import csv
import re

TE = re.compile(
    r"reverse transcriptase|transposase|transposon|transposable|retrotranspos|integrase|ribonuclease h|"
    r"rnase h|gag[- ]|gag-pre|ulp1 protease|\bhat\b|hat[, ].*dimeri|mule\b|mudr|mutator|copia|gypsy|ty1|ty3|"
    r"helitron|en/spm|cacta|pol polyprotein|polyprotein|retrovirus|retrotrans|tnt 1-94|harbinger|rrt15|art2|"
    r"far1|duf4283|duf ?659|duf4216|duf4218|pmz-type|swim-type|zinc knuckle|cchc-type|transposition|"
    r"regulator of rdna transcription protein 15|dna/rna polymerase superfamily", re.I)
SOURCES = ["InterPro", "Pfam", "PANTHER", "eggNOG", "Diamond2GO"]


def read_terms(eggnog, interproscan, diamond2go):
    """source -> (term id, name) -> set of genes."""
    terms = {s: collections.defaultdict(set) for s in SOURCES}
    for f in csv.reader(open(interproscan), delimiter="\t"):
        gene = f[0]
        if len(f) > 12 and f[11] not in ("", "-"):
            terms["InterPro"][(f[11], f[12])].add(gene)
        if f[3] == "Pfam":
            terms["Pfam"][(f[4], f[5])].add(gene)
        elif f[3] == "PANTHER":
            terms["PANTHER"][(f[4], f[5])].add(gene)
    for line in open(eggnog):
        if line.startswith("#"):
            continue
        f = line.rstrip("\n").split("\t")
        if f[7] not in ("", "-"):
            terms["eggNOG"][("", f[7])].add(f[0])
    with open(diamond2go) as fh:
        next(fh)
        for line in fh:
            f = line.rstrip("\n").split("\t")
            if len(f) < 2:
                continue
            desc = f[1].strip().split(" ", 1)[-1].strip() if " " in f[1].strip() else f[1].strip()
            if desc:
                terms["Diamond2GO"][("", desc)].add(f[0])
    return terms


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--eggnog", required=True, help="eggNOG-mapper .emapper.annotations of the main proteins")
    ap.add_argument("--interproscan", required=True, help="InterProScan TSV of the main proteins")
    ap.add_argument("--diamond2go", required=True, help="Diamond2GO TSV of the main proteins")
    ap.add_argument("--new-genes", help="IDs of genes absent from the previous annotation (one per line)")
    ap.add_argument("--min-genes", type=int, default=100, help="keep terms with more than this many genes")
    ap.add_argument("--out", required=True, help="output prefix")
    args = ap.parse_args()

    new = set(open(args.new_genes).read().split()) if args.new_genes else None
    terms = read_terms(args.eggnog, args.interproscan, args.diamond2go)
    header = ["source", "term", "name", "genes", "te_keyword"] + (["new_genes", "new_pct"] if new else [])
    rows = []
    for source in SOURCES:
        for (term, name), genes in terms[source].items():
            if len(genes) <= args.min_genes:
                continue
            m = TE.search(name)
            row = [source, term, name, len(genes), m.group(0).lower() if m else ""]
            if new is not None:
                n = len(genes & new)
                row += [n, round(100 * n / len(genes), 1)]
            rows.append(row)
    rows.sort(key=lambda r: (SOURCES.index(r[0]), -r[3]))

    with open(args.out + ".tsv", "w") as out:
        out.write("\t".join(header) + "\n")
        for r in rows:
            out.write("\t".join(str(x) for x in r) + "\n")

    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    readme = wb.active
    readme.title = "README"
    for line in __doc__.strip().splitlines():
        readme.append([line])
    readme.append([])
    readme.append([f"min_genes = {args.min_genes} ; te_keyword non vide = surligné en orange"])
    fill = PatternFill("solid", fgColor="FCE4D6")
    for title, subset in [("all", rows)] + [(s, [r for r in rows if r[0] == s]) for s in SOURCES]:
        ws = wb.create_sheet(title)
        ws.append(header + ["à supprimer (oui/non)"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in subset:
            ws.append(r)
            if r[4]:
                for c in ws[ws.max_row]:
                    c.fill = fill
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col, width in zip("ABCDEFGH", (11, 14, 70, 8, 22, 10, 9, 20)):
            ws.column_dimensions[col].width = width
    wb.save(args.out + ".xlsx")
    for s in SOURCES:
        sub = [r for r in rows if r[0] == s]
        print(f"{s}\t{len(sub)} terms\t{sum(1 for r in sub if r[4])} TE-like")


if __name__ == "__main__":
    main()
